#!/usr/bin/env python3
"""Generate PAO_starter_blank.xlsx - the blank PAO / multiple-encoding starter matrix.

Tabs: instructions | digits_0-9 | pao_00-99 | drill
Same conventions as the other matrices in this repo: Learned flag + SUM row,
'attention doublon' collision column, 'non retenu mais convient' bench column.

Usage: python3 make_template.py   (requires: pip install openpyxl)
"""
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "PAO_starter_blank.xlsx")

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
HINT_FONT = Font(italic=True, color="808080")

MAJOR = {0: "s, z", 1: "t, d", 2: "n", 3: "m", 4: "r", 5: "l",
         6: "j, ch, sh, soft g", 7: "k, hard c/g, q", 8: "f, v", 9: "p, b"}

INSTRUCTIONS = [
    (1, "read", "open README.md in this folder - it explains every concept used below (peg, Major system, PAO, multiple encoding)"),
    (2, "digits_0-9", "invent YOUR OWN shape image and rhyme image for each digit (10 minutes; the Major consonants are prefilled as reference)"),
    (3, "pao_00-99", "for each number 00-99: turn its Major consonants into a PERSON (25 = n+l = Nala, Nelly...), then give that person a signature ACTION and OBJECT"),
    (4, "collision rule", "before accepting an entry, Ctrl+F the whole sheet; if the word/person is already used, note it in 'attention doublon' (déjà pris) and pick another"),
    (5, "bench rule", "a good candidate you did not keep goes to 'non retenu mais convient' - never re-evaluate from scratch"),
    (6, "pace", "fill ONE decade (10 rows) per session; vivid beats clever - a person you can SEE, an action with movement, an object you could hold"),
    (7, "drill", "press F9 on the drill tab: a random 6-digit number appears; visualize Person(pair1) + Action(pair2) + Object(pair3) as ONE scene, then check the lookups"),
    (8, "track", "mark Learned=1 only after you recall a row 24h later without help - the SUM at the bottom is your progress bar"),
    (9, "anki (LAST)", "only when a decade is stable, export it to Anki - see 01.Millenium/tuto and 06.Programming_Millenium/build.py for the pipeline; the spreadsheet stays the source of truth"),
    ("", "", ""),
    ("example", "pao_00-99", "25: Major 2=n 5=l -> Person: Nala (Lion King). Action: licking her paw. Object: a lily."),
    ("example", "drill", "731942 -> Person of 73 + Action of 19 + Object of 42 = one single scene stored at one palace locus (6 digits, 1 image)"),
    ("why", "multiple encoding", "one number, several routes: shape (looks like), rhyme (sounds like), Major (consonants), PAO (who/does/what), palace (where) - each extra route is a backup retrieval path"),
]


def sheet_instructions(ws):
    ws["A1"], ws["B1"], ws["C1"] = "step", "where", "what to do"
    for c in ("A1", "B1", "C1"):
        ws[c].font, ws[c].fill = BOLD, HEAD_FILL
    for i, (nb, where, what) in enumerate(INSTRUCTIONS, start=2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = nb, where, what
        ws[f"C{i}"].alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 120


def sheet_digits(ws):
    headers = ["Learned (0/1)", "Digit", "Shape image (looks like...)",
               "Rhyme image (sounds like...)", "Major consonants (reference)", "Notes"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.font, c.fill = BOLD, HEAD_FILL
    for d in range(10):
        r = 2 + d
        ws.cell(r, 1, 0)
        ws.cell(r, 2, d)
        ws.cell(r, 5, MAJOR[d])
    # one greyed hint the user overwrites
    ws.cell(2, 3, "e.g. 0 = a wheel, an egg...").font = HINT_FONT
    ws["A14"] = "=SUM(A2:A11)"
    ws["B14"] = "learned / 10"
    for col, w in zip("ABCDEF", (11, 6, 30, 30, 24, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def sheet_pao(ws):
    headers = ["Learned (0/1)", "Number", "Person", "Action", "Object",
               "Attention doublon / déjà pris", "Non retenu mais convient", "Notes / Major sounds"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.font, c.fill = BOLD, HEAD_FILL
    for n in range(100):
        r = 2 + n
        ws.cell(r, 1, 0)
        ws.cell(r, 2, f"{n:02d}")
        tens, units = divmod(n, 10)
        ws.cell(r, 8, f"{MAJOR[tens].split(',')[0]} + {MAJOR[units].split(',')[0]}").font = HINT_FONT
    ws["A104"] = "=SUM(A2:A101)"
    ws["B104"] = "learned / 100"
    for col, w in zip("ABCDEFGH", (11, 8, 24, 24, 24, 30, 28, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:H101"


def sheet_drill(ws):
    ws["A1"] = "6-digit drill - press F9 (recalculate) for a new number"
    ws["A1"].font = BOLD
    ws["A3"] = "number:"
    ws["B3"] = '=TEXT(RANDBETWEEN(0,999999),"000000")'
    ws["B3"].font = Font(bold=True, size=16)
    # pair -> P/A/O lookups against the pao tab (B=number, C=Person, D=Action, E=Object)
    rows = [("Person of", 1, 2), ("Action of", 3, 3), ("Object of", 5, 4)]
    for i, (label, start, lookup_col) in enumerate(rows):
        r = 5 + i
        ws.cell(r, 1, label)
        ws.cell(r, 2, f"=MID($B$3,{start},2)")
        ws.cell(r, 3, f"=IFERROR(VLOOKUP(B{r},'pao_00-99'!$B$2:$E$101,{lookup_col},FALSE),\"(fill pao tab)\")")
    ws["A9"] = "visualize the three as ONE scene BEFORE reading column C - that scene is one palace locus holding 6 digits"
    ws["A9"].font = HINT_FONT
    for col, w in zip("ABC", (12, 12, 34)):
        ws.column_dimensions[col].width = w


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "instructions"
    sheet_instructions(ws)
    sheet_digits(wb.create_sheet("digits_0-9"))
    sheet_pao(wb.create_sheet("pao_00-99"))
    sheet_drill(wb.create_sheet("drill"))
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
